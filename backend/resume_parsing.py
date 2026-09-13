"""
Resume file handling — pure file-format text extraction, no model calls.

Turning the extracted text into structured skills/experience/education/
projects/preferred_roles is a separate step (see agent.extract_resume_data),
kept in agent.py since that's where the LLM fallback-chain infrastructure
already lives.

Requires two extra packages beyond what's already in requirements.txt:

    pip install pypdf python-docx
"""

import io

MAX_RESUME_BYTES = 5 * 1024 * 1024  # 5 MB
ALLOWED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")


class UnsupportedResumeFormat(Exception):
    """Raised for a file type we don't know how to read, or one that's too large."""


def validate_upload(filename: str, content: bytes):
    """Raises UnsupportedResumeFormat with a user-facing message if this upload can't be accepted."""
    name = (filename or "").lower().strip()

    if not name.endswith(ALLOWED_EXTENSIONS):
        raise UnsupportedResumeFormat(
            "Unsupported file type — please upload a PDF, DOCX, or plain text resume "
            f"({', '.join(ALLOWED_EXTENSIONS)})."
        )

    if len(content) > MAX_RESUME_BYTES:
        raise UnsupportedResumeFormat("That resume file is too large — the limit is 5 MB.")

    if len(content) == 0:
        raise UnsupportedResumeFormat("That file appears to be empty.")


def extract_text_from_upload(filename: str, content: bytes) -> str:
    """Extracts plain text from a resume file. Returns "" if nothing readable was found."""
    name = (filename or "").lower()

    if name.endswith(".pdf"):
        return _extract_pdf_text(content)
    if name.endswith(".docx"):
        return _extract_docx_text(content)
    # .txt, .md, or anything else that made it past validate_upload -- try
    # decoding as plain text rather than rejecting it outright.
    return _decode_best_effort(content)


def _extract_pdf_text(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages).strip()


def _extract_docx_text(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    # python-docx's `.paragraphs` skips table cells -- resumes sometimes lay
    # out skills/experience in tables, so pull those in too.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text.strip())

    return "\n".join(paragraphs).strip()


def _decode_best_effort(content: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return content.decode(encoding).strip()
        except (UnicodeDecodeError, LookupError):
            continue
    return ""